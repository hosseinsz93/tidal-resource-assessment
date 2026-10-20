"""Adapters normalize sources into one long table, without changing raw files.

Canonical columns: time, bin, depth_m, depth_reference, east_mps, north_mps,
source_file. depth_m is a positive coordinate with an explicit datum.
"""
from pathlib import Path
from datetime import datetime
import re
import numpy as np
import pandas as pd
from .config import Dataset, apply_coops_station_preset, input_file


def speed_scale(header):
    unit = str(header).lower().replace(" ", "")
    if "mm/s" in unit or "mm/sec" in unit:
        return 0.001
    if "cm/s" in unit or "cm/sec" in unit:
        return 0.01
    if "knots" in unit or "knot" in unit:
        return 0.5144444444444445
    if "m/s" in unit or "m/sec" in unit:
        return 1.0
    raise ValueError(f"Unrecognized velocity units: {header!r}. Supply explicit m/s, cm/sec, mm/s or knots.")


def rotate_heading(east, north, offset):
    angle = np.deg2rad(offset)
    return east * np.cos(angle) + north * np.sin(angle), north * np.cos(angle) - east * np.sin(angle)


def load_coops(path, settings):
    preset = apply_coops_station_preset(settings)
    paths = sorted(path.glob("*.csv")) if path.is_dir() else [path]
    pattern = re.compile(r"^([A-Za-z0-9]+)_(\d+)m(\d+)cm-.*\.csv$", re.I)
    found = []
    for file in paths:
        match = pattern.match(file.name)
        if match and match[1] == settings.site:
            found.append((file, int(match[2]) + int(match[3]) / 100))
    if not found:
        raise ValueError(f"No depth-bin CSVs for {settings.site} in {path}. Annual event predictions are not depth-bin observations.")
    depths = sorted({d for _, d in found}, reverse=settings.coops_adcp_orientation == "upward")
    geometry = (settings.water_depth_m, settings.coops_instrument_height_above_bed_m,
                settings.coops_bin1_center_distance_m, settings.coops_bin_size_m)
    if settings.coops_depth_reference == "below_surface" and all(value is not None for value in geometry):
        first_height = settings.coops_instrument_height_above_bed_m + settings.coops_bin1_center_distance_m
        bins = {depth: int(round((settings.water_depth_m - depth - first_height)
                                 / settings.coops_bin_size_m)) + 1 for depth in depths}
        if len(set(bins.values())) != len(bins) or min(bins.values()) < 1:
            raise ValueError("CO-OPS depths cannot be mapped uniquely to the configured ADCP bin geometry.")
    else:
        bins = {depth: i + 1 for i, depth in enumerate(depths)}
    frames = []
    for file, depth in found:
        raw = pd.read_csv(file)
        raw.columns = raw.columns.str.strip()
        time_col = next((c for c in raw if c.lower().startswith("date")), None)
        speed_col = next((c for c in raw if c.lower().startswith("speed")), None)
        dir_col = next((c for c in raw if c.lower().startswith("dir")), None)
        if not all((time_col, speed_col, dir_col)):
            raise ValueError(f"{file.name}: expected Date/Time, Speed (units), Dir (true).")
        if "true" not in dir_col.lower():
            raise ValueError(f"{file.name}: direction must explicitly be a true compass bearing.")
        speed = pd.to_numeric(raw[speed_col], errors="coerce") * speed_scale(speed_col)
        bearing = pd.to_numeric(raw[dir_col], errors="coerce")
        speed = speed.where((speed >= 0) & bearing.between(0, 360))
        theta = np.deg2rad(bearing)
        frames.append(pd.DataFrame({
            "time": pd.to_datetime(raw[time_col], errors="coerce", format="mixed"),
            "bin": bins[depth], "depth_m": depth,
            "depth_reference": settings.coops_depth_reference,
            "east_mps": speed * np.sin(theta), "north_mps": speed * np.cos(theta),
            "source_file": file.name,
        }))
    messages = []
    if settings.time_zone == "unspecified":
        messages.append("CO-OPS timestamps retain the downloaded clock; its timezone is unspecified. Confirm it before external time comparisons.")
    if settings.coops_depth_reference == "unknown":
        messages.append("CO-OPS depth datum is unconfirmed. Depth profiles are available; physical log/power fits require a confirmed height above bed.")
    elif settings.coops_depth_reference == "below_surface" and settings.water_depth_m is not None:
        messages.append(
            f"CO-OPS depths were converted to height above bed using the verified {settings.water_depth_m:g} m "
            "pressure-derived deployment depth."
        )
    metadata = {"source": "coops", "site": settings.site, "time_zone": settings.time_zone,
                "depth_reference": settings.coops_depth_reference,
                "adcp_orientation": settings.coops_adcp_orientation,
                "water_depth_m": settings.water_depth_m}
    if preset:
        metadata["deployment_metadata"] = {
            "nominal_water_depth_m": settings.coops_nominal_water_depth_m,
            "pressure_derived_water_depth_m": settings.water_depth_m,
            "instrument_height_above_bed_m": settings.coops_instrument_height_above_bed_m,
            "bin1_center_distance_m": settings.coops_bin1_center_distance_m,
            "bin_size_m": settings.coops_bin_size_m,
            "flood_reference_degT": settings.flood_reference_deg,
            "deployment_url": "https://api.tidesandcurrents.noaa.gov/mdapi/prod/webapi/stations/LIS1016/deployments.json",
            "bins_url": "https://api.tidesandcurrents.noaa.gov/mdapi/prod/webapi/stations/LIS1016/bins.json",
        }
    return Dataset(pd.concat(frames, ignore_index=True), metadata,
                   messages, [str(f) for f, _ in found])


def load_verdant(path, settings):
    import openpyxl
    files = sorted(path.glob("*.xlsx")) if path.is_dir() else [path]
    if len(files) != len(settings.verdant_sheets):
        raise ValueError("Provide one sheet name and heading offset per workbook, in alphabetical filename order.")
    frames, deployments = [], []
    for file, sheet, offset in zip(files, settings.verdant_sheets, settings.verdant_heading_offsets):
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        try:
            if sheet not in book.sheetnames:
                raise ValueError(f"{file.name}: sheet {sheet!r} not found. Available: {book.sheetnames}")
            rows = iter(book[sheet].iter_rows(values_only=True))
            header_rows = [next(rows) for _ in range(20)]
            labels, units = header_rows[17], header_rows[18]
            bed_row, bin_row = header_rows[11], header_rows[13]
            if "distance from bed" not in str(bed_row[8]).lower() or units[0] != "(UTC)":
                raise ValueError(f"{file.name}: unrecognized workbook layout; expected distance-from-bed metadata and UTC times.")
            starts = [j for j in range(9, len(labels)) if labels[j] == "Eas"]
            if not starts:
                raise ValueError(f"{file.name}: no depth-bin velocity blocks found.")
            records = [row for row in rows if isinstance(row[0], datetime)]
            if not records:
                raise ValueError(f"{file.name}: no dated measurements.")
            times = pd.to_datetime([r[0] for r in records], utc=True).tz_localize(None)
            for j in starts:
                if list(labels[j:j+4]) != ["Eas", "Nor", "Mag", "Dir"]:
                    raise ValueError(f"{file.name}: unexpected velocity block at column {j+1}.")
                east = pd.to_numeric(pd.Series([r[j] for r in records]), errors="coerce") * speed_scale(units[j])
                north = pd.to_numeric(pd.Series([r[j+1] for r in records]), errors="coerce") * speed_scale(units[j+1])
                east, north = rotate_heading(east, north, offset)
                frames.append(pd.DataFrame({
                    "time": times, "bin": int(bin_row[j]), "depth_m": float(bed_row[j]),
                    "depth_reference": "above_bed", "east_mps": east, "north_mps": north,
                    "source_file": file.name,
                }))
            deployments.append({"file": file.name, "sheet": sheet, "heading_offset_deg": offset})
        finally:
            book.close()
    if not frames:
        raise ValueError("No Verdant workbooks found.")
    return Dataset(pd.concat(frames, ignore_index=True),
                   {"source": "verdant", "site": settings.site, "time_zone": "UTC", "deployments": deployments},
                   ["Verdant heading offsets rotate both east and north components; heights come from workbook bed-distance metadata."],
                   [str(f) for f in files])


def load_standard_csv(path, settings):
    if path.is_dir():
        raise ValueError("Choose one standardized CSV file inside inputs/.")
    frame = pd.read_csv(path)
    required = {"time", "bin", "depth_m", "depth_reference", "east_mps", "north_mps"}
    if not required.issubset(frame):
        raise ValueError(f"Standard CSV is missing: {sorted(required - set(frame))}.")
    frame["time"] = pd.to_datetime(frame["time"], errors="coerce", format="mixed")
    if isinstance(frame["time"].dtype, pd.DatetimeTZDtype):
        frame["time"] = frame["time"].dt.tz_convert("UTC").dt.tz_localize(None)
        zone = "UTC"
    else:
        zone = settings.time_zone
    frame["source_file"] = path.name
    return Dataset(frame, {"source": "standard_csv", "site": settings.site, "time_zone": zone}, [], [str(path)])


def validate_dataset(dataset, settings):
    frame = dataset.data.copy()
    original = len(frame)
    for col in ("bin", "depth_m", "east_mps", "north_mps"):
        frame[col] = pd.to_numeric(frame[col], errors="coerce")
    frame = frame.dropna(subset=["time", "bin", "depth_m"])
    invalid_metadata = original - len(frame)
    if frame.empty:
        raise ValueError("No measurements with valid time/bin/depth metadata.")
    if ((frame.bin < 1) | (frame.bin % 1 != 0) | (frame.depth_m < 0)).any():
        raise ValueError("Bin numbers must be positive integers; depth coordinates must be nonnegative metres.")
    frame["bin"] = frame.bin.astype(int)
    if not set(frame.depth_reference).issubset({"unknown", "above_bed", "below_surface"}):
        raise ValueError("Use depth_reference = above_bed, below_surface, or unknown.")
    mapping = frame[["bin", "depth_m", "depth_reference"]].drop_duplicates()
    if mapping.bin.duplicated().any():
        raise ValueError("A bin has different depth metadata across deployments. Regrid each deployment explicitly first.")
    duplicated = frame.duplicated(["time", "bin"], keep=False)
    if duplicated.any():
        check = frame.loc[duplicated].groupby(["time", "bin"])[["east_mps", "north_mps"]].nunique(dropna=False)
        if (check > 1).any().any():
            raise ValueError("Overlapping files contain conflicting velocities for the same timestamp and bin. Resolve the overlap in the inputs.")
    n_duplicates = int(frame.duplicated(["time", "bin"]).sum())
    frame = frame.drop_duplicates(["time", "bin"])
    unknown = set(settings.excluded_bins) - set(frame.bin)
    if unknown:
        raise ValueError(f"Excluded bins are not present: {sorted(unknown)}.")
    frame = frame[~frame.bin.isin(settings.excluded_bins)].copy()
    speed = np.hypot(frame.east_mps, frame.north_mps)
    invalid = ~np.isfinite(speed) | (speed > settings.max_speed_mps)
    frame["valid_velocity"] = ~invalid
    frame.loc[invalid, ["east_mps", "north_mps"]] = np.nan
    frame["speed_mps"] = np.hypot(frame.east_mps, frame.north_mps)
    frame["direction_degT"] = np.degrees(np.arctan2(frame.east_mps, frame.north_mps)) % 360
    frame["height_above_bed_m"] = np.where(frame.depth_reference == "above_bed", frame.depth_m, np.nan)
    if settings.water_depth_m is not None:
        below = frame.depth_reference == "below_surface"
        frame.loc[below, "height_above_bed_m"] = settings.water_depth_m - frame.loc[below, "depth_m"]
        heights = frame.height_above_bed_m.dropna()
        if (heights < 0).any() or (heights > settings.water_depth_m).any():
            raise ValueError("Depth coordinates are incompatible with the specified total water depth.")
    geometry = (settings.coops_instrument_height_above_bed_m,
                settings.coops_bin1_center_distance_m, settings.coops_bin_size_m)
    if settings.source == "coops" and all(value is not None for value in geometry):
        bin_geometry = frame[["bin", "height_above_bed_m"]].drop_duplicates().sort_values("bin")
        expected = (settings.coops_instrument_height_above_bed_m
                    + settings.coops_bin1_center_distance_m
                    + settings.coops_bin_size_m * (bin_geometry.bin.to_numpy() - 1))
        difference = np.abs(bin_geometry.height_above_bed_m.to_numpy() - expected)
        if not np.isfinite(difference).all() or (difference > 0.05).any():
            raise ValueError("CO-OPS bin depths disagree with the configured ADCP orientation or deployment geometry.")
        dataset.metadata["height_above_bed_range_m"] = [
            float(bin_geometry.height_above_bed_m.min()),
            float(bin_geometry.height_above_bed_m.max()),
        ]
        dataset.metadata["height_geometry_max_difference_m"] = float(difference.max())
    if frame.empty or not frame.valid_velocity.any():
        raise ValueError("No valid velocities remain after bin selection and quality control.")
    times = frame.time.drop_duplicates().sort_values()
    if len(times) < 3:
        raise ValueError("At least three distinct timestamps are needed.")
    dataset.metadata["quality_control"] = {
        "input_rows": original, "rows_without_time_bin_depth": invalid_metadata,
        "identical_duplicates_removed": n_duplicates, "invalid_velocity_rows": int(invalid.sum()),
        "retained_rows": len(frame), "retained_bins": int(frame.bin.nunique()),
    }
    if invalid.any():
        dataset.messages.append(f"{int(invalid.sum()):,} velocity rows were missing or exceeded the speed limit; retained as missing, never as slack.")
    dataset.data = frame.sort_values(["time", "bin"]).reset_index(drop=True)
    return dataset


def load(root, settings):
    settings.validate()
    path = input_file(root, settings.input_path)
    loader = {"coops": load_coops, "verdant": load_verdant, "standard_csv": load_standard_csv}[settings.source]
    return validate_dataset(loader(path, settings), settings)
