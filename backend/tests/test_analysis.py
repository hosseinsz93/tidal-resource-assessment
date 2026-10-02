"""Numerical regression checks independent of legacy notebook outputs."""
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
from backend.tidal import Settings, preprocess, process
from backend.tidal.config import Dataset
from backend.tidal.loaders import validate_dataset, rotate_heading, speed_scale, load_coops, load_verdant
from backend.tidal.analysis import time_support, estimate_axis, energy, tidal_cycles, profiles


def synthetic(tmp_path, bins=4):
    times = pd.date_range("2024-01-01", periods=600, freq="6min")
    wave = np.sin(np.arange(len(times))*2*np.pi/120)
    frame = pd.concat([pd.DataFrame({"time": times, "bin": b, "depth_m": b,
                                    "depth_reference": "above_bed", "east_mps": -wave*(b/bins)**0.2,
                                    "north_mps": 0.0}) for b in range(1, bins+1)], ignore_index=True)
    path = tmp_path / "inputs" / "custom"
    path.mkdir(parents=True)
    frame.to_csv(path / "test.csv", index=False)
    cfg = Settings(source="standard_csv", site="Test", input_path="inputs/custom/test.csv",
                   flood_reference_deg=270, energy_enabled=False, hub_bin=1, water_depth_m=5)
    return frame, cfg


@pytest.mark.parametrize("unit", ["ns", "us", "ms"])
def test_time_support_does_not_count_gaps_or_depend_on_storage_unit(unit):
    times = pd.DatetimeIndex(["2024-01-01 00:00", "2024-01-01 00:06", "2024-01-01 02:00", "2024-01-01 02:06"]).as_unit(unit)
    weights, blocks, nominal = time_support(times)
    assert nominal == 360
    np.testing.assert_array_equal(weights, [360]*4)
    np.testing.assert_array_equal(blocks, [0, 0, 1, 1])


def test_shared_labels_and_harmonic_fit(tmp_path):
    _, cfg = synthetic(tmp_path)
    assessment = process(preprocess(tmp_path, cfg))
    avg, resolved = assessment.tables["depth_averaged"], assessment.tables["depth_resolved"]
    assert assessment.bearing == pytest.approx(270)
    assert (avg.loc[avg.phase == "FLOOD", "east_mps"] < 0).all()
    expected = avg.set_index("time").phase
    assert (resolved.phase == resolved.time.map(expected)).all()
    fit = assessment.tables["cycle_fit"].iloc[0]
    assert fit.r_squared > 0.999
    assert assessment.tables["cycles"].duration_hours.between(11.99, 12.01).all()
    good_power = assessment.tables["profile_fits"].query("model == 'power' and status == 'ok'")
    np.testing.assert_allclose(good_power.alpha, 5, rtol=1e-6)


def test_missing_bin_does_not_turn_into_slack(tmp_path):
    frame, cfg = synthetic(tmp_path)
    frame.loc[(frame.bin == 1) & (frame.time == frame.time.iloc[20]), "east_mps"] = np.nan
    frame.to_csv(tmp_path / cfg.input_path, index=False)
    result = process(preprocess(tmp_path, cfg))
    row = result.tables["depth_averaged"].iloc[20]
    assert row.phase == "MISSING"
    assert row.bin_coverage == 0.75


def test_coops_any_number_of_segments_and_units(tmp_path):
    for i in range(3):
        frame = pd.DataFrame({"Date/Time": pd.date_range("2024-01-01", periods=9, freq="6min")[i*3:i*3+3],
                              " Speed (cm/sec)": [100]*3, " Dir (true)": [90]*3})
        frame.to_csv(tmp_path/f"TEST_05m76cm-2024-01-0{i+1}.csv", index=False)
    dataset = load_coops(tmp_path, Settings(site="TEST"))
    assert len(dataset.data) == 9
    assert dataset.data.depth_m.iloc[0] == 5.76
    np.testing.assert_allclose(dataset.data.east_mps, 1)
    np.testing.assert_allclose(dataset.data.north_mps, 0, atol=1e-12)


def test_conflicting_overlap_rejected(tmp_path):
    frame, cfg = synthetic(tmp_path)
    bad = frame.iloc[[0]].copy()
    bad["east_mps"] = 20
    dataset = Dataset(pd.concat([frame, bad]))
    with pytest.raises(ValueError, match="conflicting"):
        validate_dataset(dataset, cfg)


def test_heading_rotation_preserves_speed_and_units():
    east, north = rotate_heading(np.array([1.0]), np.array([0.0]), 90)
    assert east[0] == pytest.approx(0, abs=1e-12)
    assert north[0] == pytest.approx(-1)
    assert np.hypot(east, north)[0] == pytest.approx(1)
    assert speed_scale("Speed (mm/s)") == .001
    assert speed_scale("Speed (cm/sec)") == .01
    with pytest.raises(ValueError, match="units"):
        speed_scale("Speed")


def test_energy_time_weights_and_cutoffs():
    resolved = pd.DataFrame({"time": pd.date_range("2024-01-01", periods=4, freq="6min"), "bin": 1,
                             "along_mps": [0.2, 1.0, -1.0, 3.5], "weight_s": [360, 360, 720, 360],
                             "phase": ["SLACK", "FLOOD", "EBB", "FLOOD"]})
    curve = pd.DataFrame({"speed_mps": [0, 4], "power_kw": [100, 100]})
    bins, summary = energy(resolved, curve, Settings(hub_bin=1, rated_kw=100))
    assert bins.probability.sum() == pytest.approx(1)
    assert bins.filter(like="prob_").sum().sum() == pytest.approx(1)
    assert summary.mean_power_kw.iloc[0] == pytest.approx(60)
    assert summary.annual_kwh_per_turbine.iloc[0] == pytest.approx(8760*60)
    assert bins.loc[bins.speed_center_mps >= 3.2, "power_kw"].eq(0).all()


def test_cycles_never_cross_a_gap(tmp_path):
    _, cfg = synthetic(tmp_path)
    result = process(preprocess(tmp_path, cfg))
    avg = result.tables["depth_averaged"].copy()
    avg.loc[avg.index >= 240, "time"] += pd.Timedelta(days=5)
    avg.loc[avg.index >= 240, "analysis_block"] += 1
    cycles, _, _, _ = tidal_cycles(avg, cfg)
    assert cycles.duration_hours.max() < 13
    gap_start, gap_end = avg.time.iloc[239], avg.time.iloc[240]
    assert not ((cycles.start < gap_start) & (cycles.end > gap_end)).any()


def test_unknown_depth_does_not_trigger_inferred_fits(tmp_path):
    frame, cfg = synthetic(tmp_path)
    frame["depth_reference"] = "unknown"
    frame.to_csv(tmp_path/cfg.input_path, index=False)
    result = process(preprocess(tmp_path, cfg))
    assert result.tables["profile_fits"].status.ne("ok").all()


def test_verdant_adapter_reads_geometry_and_rotates_vectors(tmp_path):
    import openpyxl
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Station B"
    sheet.cell(12, 9, "Distance from bed (m):")
    for j, value in enumerate(["Eas", "Nor", "Mag", "Dir"], 10):
        sheet.cell(18, j, value)
        sheet.cell(19, j, "(mm/s)" if value != "Dir" else "(deg T)")
        sheet.cell(12, j, 1.67)
        sheet.cell(14, j, 1)
    sheet.cell(19, 1, "(UTC)")
    for row, time in enumerate(pd.date_range("2024-01-01", periods=3, freq="10min"), 21):
        sheet.cell(row, 1, time.to_pydatetime())
        sheet.cell(row, 10, 1000)
        sheet.cell(row, 11, 0)
    book.save(tmp_path/"deployment.xlsx")
    cfg = Settings(verdant_sheets=("Station B",), verdant_heading_offsets=(90,))
    data = load_verdant(tmp_path, cfg).data
    np.testing.assert_allclose(data.north_mps, -1)
    assert data.depth_m.eq(1.67).all()
    assert data.depth_reference.eq("above_bed").all()


def test_notebook_presets_and_run_log_restore(tmp_path):
    import json
    from backend.tidal.notebook import SetupPanel
    panel = SetupPanel(tmp_path)
    panel.controls["source"].value = "verdant"
    cfg = panel.get_settings()
    assert cfg.excluded_bins == (45, 46, 47, 48, 49)
    assert cfg.time_zone == "UTC"
    restored = Settings(**json.loads(json.dumps(cfg.to_dict())))
    restored_panel = SetupPanel(tmp_path, restored)
    assert restored_panel.get_settings() == cfg
    panel.controls["source"].value = "coops"
    assert panel.get_settings().excluded_bins == ()


def test_settings_snapshot_and_low_coverage_guard(tmp_path):
    frame, cfg = synthetic(tmp_path)
    assessment = preprocess(tmp_path, cfg)
    cfg.slack_mps = 10
    assert assessment.settings.slack_mps == .4
    frame.loc[frame.bin == 1, "east_mps"] = np.nan
    frame.to_csv(tmp_path/cfg.input_path, index=False)
    with pytest.raises(ValueError, match="coverage"):
        process(preprocess(tmp_path, cfg))


def test_physical_log_fit_uses_true_height_and_near_bed_region():
    z = np.array([.5, 1., 1.5, 5., 8.])
    expected_ustar, expected_z0 = .12, .01
    u = expected_ustar/.41 * np.log(z/expected_z0)
    resolved = pd.DataFrame({"time": pd.Timestamp("2024-01-01"), "bin": np.arange(1, 6),
                             "depth_m": z, "height_above_bed_m": z, "depth_reference": "above_bed",
                             "phase": "FLOOD", "along_mps": u, "speed_mps": u, "event_id": 0})
    _, _, fits, _ = profiles(resolved, Settings(water_depth_m=10))
    fit = fits[fits.model == "log"].iloc[0]
    assert fit.status == "ok"
    assert fit.fit_points == 3
    assert fit.u_star_mps == pytest.approx(expected_ustar)
    assert fit.z0_m == pytest.approx(expected_z0)
